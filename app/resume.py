"""Turn an uploaded resume (PDF, DOCX, TXT, MD) into plain text plus facts."""
import io
import re

from . import extract

MIN_CHARS = 200
UNREADABLE = "Couldn't read enough text from this file. If it's a scanned PDF, export it as a text PDF or DOCX."


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".docx"):
        import docx
        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Upload a PDF, DOCX, TXT or MD file.")
    return _tidy(text)


def _tidy(text: str) -> str:
    text = re.sub(r"[ \t\xa0]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def read_resume(filename: str, data: bytes) -> tuple[str, dict | None]:
    """Return (text, facts). Facts come from Gemini and are None without it.
    A scanned PDF that has no text layer is transcribed by Gemini instead."""
    text = extract_text(filename, data)  # raises ValueError for unsupported types
    local_ok = len(text) >= MIN_CHARS
    facts = extract.resume_facts(filename, data, text if local_ok else None)
    if not local_ok:
        text = _tidy((facts or {}).pop("plain_text", "") or "")
        if len(text) < MIN_CHARS:
            raise ValueError(UNREADABLE)
    return text, facts
